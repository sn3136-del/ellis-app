"""The workbook shows the two documented labels where the record surface
records a documented blank, and a filled cell verbatim."""
import asyncio
import io

from openpyxl import load_workbook

from app import main
from app.visa_snapshot import tstation


def _workbook(rows, monkeypatch):
    monkeypatch.setattr(main, '_tstation_rows', lambda *a, **kw: rows)
    monkeypatch.setattr(main, 'require_quality_control', lambda p: None)
    response = main.travel_database_export(db=None, p=None)

    async def consume():
        return b''.join([chunk async for chunk in response.body_iterator])

    return load_workbook(io.BytesIO(asyncio.run(consume())), read_only=True)


def test_export_labels_documented_blanks_and_keeps_filled_cells(monkeypatch):
    # A visa-free record: the validity cells cannot apply, the destination
    # was checked for a policy end date and does not publish one, and the
    # stay is a filled value.
    row = {'travel_document_type': 'ordinary_passport', 'travel_document_country': 'JPN',
           'destination_country': 'KOR', 'travel_purpose': 'tourism',
           'visa_requirement': 'Visa-free', 'visa_requirement_detail': 'Unconditional Visa-free',
           'visa_type_name': 'No visa needed', 'max_stay_duration': 90, 'max_stay_unit': 'Day',
           'visa_fee_amount': 0, 'required_documents': 'Valid passport',
           'data_source': 'Operator', 'source_url': 'https://www.mofa.go.kr/',
           'collected_at': '2026-09-09', 'confidence_level': 'High',
           '_unpublished': ['info_validity']}
    statuses = tstation.field_status(row)
    assert statuses['info_validity'] == 'not-published'
    assert statuses['application_method'] == 'not-applicable'
    assert statuses['max_stay_duration'] == 'filled'
    workbook = _workbook([row], monkeypatch)
    header, record = list(workbook['Data'].values)[:2]
    assert header == tstation.FIELD_ORDER
    cell = dict(zip(header, record, strict=True))
    assert cell['info_validity'] == 'Not publicly available'
    assert cell['application_method'] == 'Not applicable'
    assert cell['processing_min_days'] == 'Not applicable'
    assert cell['max_stay_duration'] == 90 and cell['max_stay_unit'] == 'Day'
    assert cell['source_url'] == 'https://www.mofa.go.kr/' and cell['visa_fee_amount'] == 0
    # A gap is never blank and never invented: the owner's two labels only.
    assert cell['consulate_district'] == {'not-applicable': 'Not applicable', 'optional-empty': 'Not publicly available'}[statuses['consulate_district']]
    assert cell['entry_requirements'] == 'Not publicly available' and statuses['entry_requirements'] == 'optional-empty'
    workbook.close()


def test_export_values_match_the_record_surface_verdicts():
    row = {'visa_requirement': 'Visa Required in Advance', 'visa_type_name': 'Tourist visa',
           'visa_fee_amount': None, '_unpublished': ['visa_fee_amount', 'visa_fee_currency']}
    values = dict(zip(tstation.FIELD_ORDER, tstation.export_values(row), strict=True))
    assert values['visa_fee_amount'] == values['visa_fee_currency'] == tstation.NOT_PUBLICLY_AVAILABLE
    assert values['visa_type_name'] == 'Tourist visa'
    assert values['validity_duration'] == tstation.NOT_PUBLICLY_AVAILABLE  # a gap reads as the label, and stays a gap for grading
    assert tstation.field_status(row)['validity_duration'] == 'missing'


def test_a_stay_stated_in_words_is_exported_verbatim_not_labelled():
    row = {'visa_requirement': 'Visa Required in Advance', 'visa_type_name': 'Tourist visa',
           'max_stay_text': 'Stay is determined by the e-Pass issued on arrival'}
    values = dict(zip(tstation.FIELD_ORDER, tstation.export_values(row), strict=True))
    assert values['max_stay_duration'] == 'Stay is determined by the e-Pass issued on arrival'
    assert values['max_stay_unit'] == tstation.NOT_APPLICABLE
    # The wording is the value the record holds, so it is filled for grading
    # too, and the unit cell is the record's own "Not applicable" verdict.
    status = tstation.field_status(row)
    assert status['max_stay_duration'] == 'filled'
    assert status['max_stay_unit'] == 'not-applicable'
    assert 'max_stay' not in [f for f, v in status.items() if v == 'missing']


def test_a_stay_in_words_does_not_count_as_a_gap_for_completeness():
    row = {'visa_requirement': 'Visa Required in Advance', 'visa_requirement_detail': 'eVisa',
           'visa_type_name': 'Standard Visitor', 'validity_duration': 6, 'validity_unit': 'Month',
           'entries': 'Multiple', 'visa_fee_amount': 135, 'visa_fee_currency': 'GBP',
           'application_method': 'Online Application', 'required_documents': 'Passport',
           'consulate_district': 'Any', 'entry_requirements': 'x', 'special_conditions': 'y',
           'data_source': 'Ellis AI official-source field review', 'source_url': 'https://www.gov.uk/standard-visitor',
           'collected_at': '2026-09-11', 'info_validity': '2027-09-11', 'confidence_level': 'High',
           'travel_document_type': 'ordinary_passport', 'travel_document_country': 'RUS',
           'destination_country': 'GBR', 'travel_purpose': 'tourism',
           'max_stay_duration': None, 'max_stay_unit': None,
           'max_stay_text': 'Up to 6 calendar months per visit'}
    assert tstation.completeness(row) == 1.0
    numeric = dict(row, max_stay_duration=180, max_stay_unit='Day', max_stay_text=None)
    assert tstation.field_status(numeric)['max_stay_duration'] == 'filled'
    assert tstation.field_status(numeric)['max_stay_unit'] == 'filled'
    values = dict(zip(tstation.FIELD_ORDER, tstation.export_values(numeric), strict=True))
    assert values['max_stay_duration'] == 180 and values['max_stay_unit'] == 'Day'


def test_a_validity_stated_in_words_is_a_filled_value_with_no_unit():
    row = {'visa_requirement': 'Visa Required in Advance', 'visa_requirement_detail': 'Paper Visa',
           'visa_type_name': 'Ordinary tourist visa', 'validity_duration': None, 'validity_unit': None,
           'validity_text': 'Up to 3 months for a single or double entry visa, up to 6 months for a multiple entry visa'}
    status = tstation.field_status(row)
    assert status['validity_duration'] == 'filled' and status['validity_unit'] == 'not-applicable'
    values = dict(zip(tstation.FIELD_ORDER, tstation.export_values(row), strict=True))
    assert values['validity_duration'] == row['validity_text']
    assert values['validity_unit'] == tstation.NOT_APPLICABLE
    # Wording that only points elsewhere is not a value of its own.
    for pointer in ('As above', 'Same as the single-entry visa', 'See notes below', 'n/a'):
        gap = dict(row, validity_text=pointer)
        assert tstation.field_status(gap)['validity_duration'] == 'missing'
        assert dict(zip(tstation.FIELD_ORDER, tstation.export_values(gap), strict=True))['validity_duration'] == tstation.NOT_PUBLICLY_AVAILABLE
    # A numeric validity is untouched.
    numeric = dict(row, validity_duration=6, validity_unit='Month', validity_text=None)
    assert tstation.field_status(numeric)['validity_unit'] == 'filled'


def test_wording_predicates_follow_the_owner_rule():
    # Pointers and placeholders are gaps on both cells.
    for pointer in ('As above', 'Same as the single-entry visa', 'See notes below', 'n/a', 'Unknown', 'TBD', '-'):
        stay = {'visa_requirement': 'Visa Required in Advance', 'max_stay_text': pointer}
        assert tstation.field_status(stay)['max_stay_duration'] == 'missing', pointer
        val = {'visa_requirement': 'Visa Required in Advance', 'validity_text': pointer}
        assert tstation.field_status(val)['validity_duration'] == 'missing', pointer
    # Values the source states, including rule wording and short CJK
    # validities, are filled.
    for value in ('As decided by the consulate (minimum 15 days)', '12 months or above', '3个月一次有效',
                  '1年多次有效', 'as granted', 'Granted for the period applied for',
                  'Determined by consular officials and the visa issued.', 'Limited to the intended trip dates'):
        val = {'visa_requirement': 'Visa Required in Advance', 'validity_text': value}
        assert tstation.field_status(val)['validity_duration'] == 'filled', value
        assert tstation.field_status(val)['validity_unit'] == 'not-applicable', value
    # Wording that asserts an absence is the documented label, not a value.
    for absent in ('not published', 'Not published by the Ministry', 'Not published as a fixed number, the period of stay is shown on the visit pass'):
        val = {'visa_requirement': 'Visa Required in Advance', 'validity_text': absent}
        st = tstation.field_status(val)
        assert st['validity_duration'] == 'not-published' and st['validity_unit'] == 'not-published', absent
        cells = dict(zip(tstation.FIELD_ORDER, tstation.export_values(val), strict=True))
        assert cells['validity_duration'] == tstation.NOT_PUBLICLY_AVAILABLE
        stay = {'visa_requirement': 'Visa Required in Advance', 'max_stay_text': absent}
        assert tstation.field_status(stay)['max_stay_duration'] == 'not-published'
    val = {'visa_requirement': 'Visa Required in Advance', 'validity_text': 'Not applicable, no visa issued'}
    assert tstation.field_status(val)['validity_duration'] == 'not-applicable'
    assert dict(zip(tstation.FIELD_ORDER, tstation.export_values(val), strict=True))['validity_duration'] == tstation.NOT_APPLICABLE
    # A processing unit is not a validity.
    val = {'visa_requirement': 'Visa Required in Advance', 'validity_text': '60 Working Day'}
    assert tstation.field_status(val)['validity_duration'] == 'missing'


def test_a_visa_free_row_keeps_not_applicable_beside_stay_wording():
    row = {'visa_requirement': 'Visa-free', 'visa_requirement_detail': 'Conditional Visa-free',
           'max_stay_text': 'Up to 3 months at a time', 'validity_text': 'Up to 3 months at a time',
           'validity_duration': None, 'validity_unit': None, 'max_stay_duration': None, 'max_stay_unit': None}
    st = tstation.field_status(row)
    assert st['validity_duration'] == 'not-applicable' and st['validity_unit'] == 'not-applicable'
    cells = dict(zip(tstation.FIELD_ORDER, tstation.export_values(row), strict=True))
    assert cells['validity_duration'] == tstation.NOT_APPLICABLE
    assert cells['max_stay_duration'] == 'Up to 3 months at a time' and cells['max_stay_unit'] == tstation.NOT_APPLICABLE
    # The strip helper clears the wording too.
    stripped = tstation._strip_visa_only_fields(dict(row, visa_requirement='Visa-free'))
    assert stripped['validity_text'] is None
    # A documented not-published cell keeps its label even with wording beside it.
    doc = {'visa_requirement': 'Visa Required in Advance', 'validity_text': '6 months', '_unpublished': ['validity_duration', 'validity_unit']}
    cells = dict(zip(tstation.FIELD_ORDER, tstation.export_values(doc), strict=True))
    assert cells['validity_duration'] == tstation.NOT_PUBLICLY_AVAILABLE


def test_acceptance_summary_counts_wording_cells_as_present():
    row = {f: 'x' for f in tstation.FIELD_ORDER}
    row.update({'visa_requirement': 'Visa Required in Advance', 'validity_duration': None, 'validity_unit': None,
                'validity_text': 'Up to 3 months for a single or double entry visa, up to 6 months for a multiple entry visa',
                'max_stay_duration': None, 'max_stay_unit': None, 'max_stay_text': 'Up to 90 days per visit',
                'confidence_level': 'High', '_publication_state': 'published'})
    summary = tstation.acceptance_summary([row])
    # The wording cells count as filled; the unit cells beside them are the
    # record's own Not applicable label, so the record is documented
    # complete (the owner's approved definition) while the literal non-null
    # diagnostic still shows the two label cells.
    assert summary['documented_complete_records'] == 1
    assert summary['filled_cells'] == len(tstation.CONTRACT_FIELDS) - 2
    assert summary['complete_records'] == 0
    assert tstation.completeness(row) == 1.0
