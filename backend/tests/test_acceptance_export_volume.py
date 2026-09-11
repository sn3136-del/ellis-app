"""Exercise the actual exporter above the acceptance standard's 10k floor."""
import asyncio
import io

from openpyxl import load_workbook

from app import main
from app.visa_snapshot import tstation


def test_export_preserves_10001_rows_sources_zero_fees_and_dictionary(monkeypatch):
    count = 10001
    rows = [dict(visa_type_name=f'Product {i}', visa_fee_amount=0,
                 source_url=f'https://example.gov/product/{i}', confidence_level='Low')
            for i in range(count)]
    monkeypatch.setattr(main, '_tstation_rows', lambda *a, **kw: rows)
    monkeypatch.setattr(main, 'require_quality_control', lambda p: None)
    response = main.travel_database_export(db=None, p=None)

    async def consume():
        return b''.join([chunk async for chunk in response.body_iterator])

    workbook = load_workbook(io.BytesIO(asyncio.run(consume())), read_only=True)
    assert workbook.sheetnames == ['Data', 'Field descriptions']
    data = list(workbook['Data'].values)
    assert data[0] == tstation.FIELD_ORDER
    assert len(data) == count + 1
    name = tstation.FIELD_ORDER.index('visa_type_name')
    url = tstation.FIELD_ORDER.index('source_url')
    fee = tstation.FIELD_ORDER.index('visa_fee_amount')
    for i, record in enumerate(data[1:]):
        assert record[name] == f'Product {i}'
        assert record[url] == f'https://example.gov/product/{i}'
        assert record[fee] == 0
    descriptions = list(workbook['Field descriptions'].values)
    assert descriptions[0][0] == 'Snapshot (UTC)'
    assert descriptions[0][1].endswith('Z')
    fields = descriptions[2:2 + len(tstation.FIELD_ORDER)]
    assert {row[1] for row in fields} == set(tstation.FIELD_ORDER)
    # The two documented labels are defined on the same sheet.
    labels = {row[0] for row in descriptions[2 + len(tstation.FIELD_ORDER):] if row and row[0]}
    assert {tstation.NOT_PUBLICLY_AVAILABLE, tstation.NOT_APPLICABLE} <= labels
    workbook.close()
